"""Run what-if experiments: modify inputs in the source workbook, recalc in LibreOffice, read outputs."""
import datetime as dt
import json
import shutil
import subprocess
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).parent
SRC = HERE / "excel_labs_src.xlsx"

experiments = [
    {"id": "E2.1", "lab": 2, "title": "Lower the % threshold to 3%",
     "changes": {"BvA!B2": 0.03}, "watch": ["BvA!H6", "BvA!H7", "BvA!H8", "BvA!H9", "BvA!H10", "BvA!H11", "BvA!H12"]},
    {"id": "E2.2", "lab": 2, "title": "Switch the report period to 2026-07 (no budget loaded)",
     "changes": {"BvA!B1": "2026-07"}, "watch": ["BvA!C6", "BvA!D6", "BvA!F6", "BvA!G6", "BvA!H6", "BvA!D14"]},
    {"id": "E3.1", "lab": 3, "title": "Raise the flux % threshold to 20%",
     "changes": {"Flux!B4": 0.20}, "watch": ["Flux!I8", "Flux!I9", "Flux!I10", "Flux!I11", "Flux!I12", "Flux!I13"]},
    {"id": "E4.1", "lab": 4, "title": "Pro sells at budget price (E6 = 120)",
     "changes": {"PVM!E6": 120}, "watch": ["PVM!J8", "PVM!K8", "PVM!L8", "PVM!M8", "PVM!N8"]},
    {"id": "E4.2", "lab": 4, "title": "Same total units, budget mix (D5=994, D6=497, D7=99)",
     "changes": {"PVM!D5": 994, "PVM!D6": 497, "PVM!D7": 99}, "watch": ["PVM!K8", "PVM!L8", "PVM!N8"]},
    {"id": "E5.1", "lab": 5, "title": "Negotiate freight rate down to 11.50",
     "changes": {"CostVar!B6": 11.5}, "watch": ["CostVar!B10", "CostVar!C10", "CostVar!B12", "CostVar!C12"]},
    {"id": "E7.1", "lab": 7, "title": "Tighten amount tolerance to 1.00",
     "changes": {"BankRec!B2": 1}, "watch": ["Book!E4", "Bank!E4", "BankRec!B10", "BankRec!B15", "BankRec!B16",
                                            "BankRec!B12", "BankRec!B17", "BankRec!B19"]},
    {"id": "E7.2", "lab": 7, "title": "Bank statement balance keyed wrong (10,100)",
     "changes": {"BankRec!B5": 10100}, "watch": ["BankRec!B19", "BankRec!B20"]},
    {"id": "E8.1", "lab": 8, "title": "Roll the as-of date to 2026-09-30",
     "changes": {"AR!H2": dt.date(2026, 9, 30)},
     "watch": ["ARSummary!B4", "ARSummary!B5", "ARSummary!B6", "ARSummary!B7", "ARSummary!B8", "ARSummary!E9"]},
    {"id": "E9.1", "lab": 9, "title": "Insert an outlier: Apr-24 sales = 250,000",
     "changes": {"Forecast!C8": 250000}, "watch": ["Forecast!L19", "Forecast!M8", "Forecast!E41", "Forecast!B54"]},
    {"id": "E10.1", "lab": 10, "title": "Switch to the Upside scenario (C8 = 3)",
     "changes": {"Drivers!C8": 3}, "watch": ["Drivers!C9", "Drivers!G17", "Drivers!H18", "Checks!B9"]},
    {"id": "E10.2", "lab": 10, "title": "Base scenario with churn 3% (C5 = 0.03)",
     "changes": {"Drivers!C5": 0.03}, "watch": ["Drivers!H18", "Drivers!B33", "Drivers!D33"]},
]

results = []
work = HERE / "exp"; shutil.rmtree(work, ignore_errors=True); work.mkdir()
for ex in experiments:
    wb = load_workbook(SRC)
    for ref, v in ex["changes"].items():
        s, a = ref.split("!"); wb[s][a] = v
    p = work / f"{ex['id']}.xlsx"; wb.save(p)
    outdir = work / ex["id"]; outdir.mkdir()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(p)],
                   check=True, capture_output=True, timeout=180)
    cw = load_workbook(outdir / p.name, data_only=True)
    base = load_workbook(HERE / "calc" / "excel_labs_src.xlsx", data_only=True)
    obs = []
    for ref in ex["watch"]:
        s, a = ref.split("!")
        obs.append({"cell": ref, "before": base[s][a].value, "after": cw[s][a].value})
    ex["results"] = obs
    results.append(ex)
    print(ex["id"], ex["title"])
    for o in obs:
        print("   ", o["cell"], o["before"], "->", o["after"])
(HERE / "experiments.json").write_text(json.dumps(results, indent=1, default=str))
