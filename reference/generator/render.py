"""Render the markdown guide from template.md, filling placeholders from the verified workbook."""
import json
import re
from pathlib import Path
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

HERE = Path(__file__).parent
spec = json.loads((HERE / "spec.json").read_text())
vals = json.loads((HERE / "values.json").read_text())
exps = json.loads((HERE / "experiments.json").read_text())


from decimal import Decimal, ROUND_HALF_UP


def rhu(x, dec):
    """Round half away from zero, like Excel's display rounding."""
    q = Decimal(1).scaleb(-dec)
    return float(Decimal(repr(x)).quantize(q, rounding=ROUND_HALF_UP))


def fmt_value(v, fmt):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, str):
        if re.match(r"^\d{4}-\d{2}-\d{2} 00:00:00$", v):
            d = v[:10]
            if fmt and "mmm" in fmt:
                import datetime as dt
                return dt.date.fromisoformat(d).strftime("%b-%y")
            return d
        return v
    if fmt:
        if "%" in fmt:
            dec = len(fmt.split(".")[1].split("%")[0]) if "." in fmt else 0
            s = f"{rhu(v*100, dec):,.{dec}f}%"
            return ("+" + s if v > 0 and fmt.startswith("+") else s)
        if "." in fmt:
            dec = len(fmt.split(";")[0].split(".")[1])
            s = f"{rhu(abs(v), dec):,.{dec}f}"
        else:
            s = f"{rhu(abs(v), 0):,.0f}"
        if v < 0:
            return f"({s})" if ";(" in fmt else "-" + s
        return s
    if isinstance(v, float) and not v.is_integer():
        return f"{v:,.4f}".rstrip("0").rstrip(".")
    if isinstance(v, (int, float)):
        return f"{v:,.0f}" if abs(v) >= 1000 else f"{v:g}"
    return str(v)


def cell_fmt(sheet, addr):
    return spec["sheets"][sheet].get(addr, {}).get("fmt")


def show(sheet, addr):
    return fmt_value(vals[sheet].get(addr), cell_fmt(sheet, addr))


def values_grid(ref):
    sheet, rng = ref.split("!")
    c1, r1, c2, r2 = range_boundaries(rng)
    cols = [get_column_letter(c) for c in range(c1, c2 + 1)]
    out = ["| | " + " | ".join(cols) + " |", "|---|" + "---|" * len(cols)]
    for r in range(r1, r2 + 1):
        cells = []
        for c in cols:
            txt = show(sheet, f"{c}{r}").replace("|", "\\|")
            cells.append(txt)
        out.append(f"| **{r}** | " + " | ".join(cells) + " |")
    return "\n".join(out)


def formulas_table(ref):
    sheet, rng = ref.split("!")
    addrs = []
    for part in rng.split(","):
        if ":" in part:
            c1, r1, c2, r2 = range_boundaries(part)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    addrs.append(f"{get_column_letter(c)}{r}")
        else:
            addrs.append(part)
    out = ["| Cell | Formula | Result |", "|---|---|---|"]
    for a in addrs:
        cell = spec["sheets"][sheet].get(a)
        if not cell or "f" not in cell:
            continue
        out.append(f"| `{sheet}!{a}` | `{cell['f']}` | {show(sheet, a)} |")
    return "\n".join(out)


def experiments_table(lab):
    blocks = []
    for ex in exps:
        if ex["lab"] != int(lab):
            continue
        changes = ", ".join(f"`{k}` → `{v}`" for k, v in ex["changes"].items())
        lines = [f"**{ex['id']} — {ex['title']}**  ", f"Change: {changes}", "",
                 "| Watch cell | Before | After |", "|---|---|---|"]
        for o in ex["results"]:
            s, a = o["cell"].split("!")
            f = cell_fmt(s, a)
            lines.append(f"| `{o['cell']}` | {fmt_value(o['before'], f)} | {fmt_value(o['after'], f)} |")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def full_spec():
    parts = []
    for name in spec["sheetOrder"]:
        cells = spec["sheets"][name]
        lines = [f"#### Sheet `{name}` ({len(cells)} cells)", "", "```json", "{"]
        items = list(cells.items())
        for i, (a, c) in enumerate(items):
            comma = "," if i < len(items) - 1 else ""
            lines.append(f'  "{a}": {json.dumps(c)}{comma}')
        lines += ["}", "```", ""]
        parts.append("\n".join(lines))
    return "\n".join(parts)


def assertions():
    lines = ["```json", "["]
    a = spec["assertions"]
    for i, x in enumerate(a):
        comma = "," if i < len(a) - 1 else ""
        lines.append("  " + json.dumps(x) + comma)
    lines += ["]", "```"]
    return "\n".join(lines)


def experiments_json():
    slim = [{"id": e["id"], "lab": e["lab"], "title": e["title"], "changes":
             {k: (v if not isinstance(v, str) or not v.endswith("00:00:00") else v[:10]) for k, v in e["changes"].items()},
             "expect": {o["cell"]: o["after"] for o in e["results"]}} for e in exps]
    return "```json\n" + json.dumps(slim, indent=1, default=str) + "\n```"


tpl = (HERE / "template.md").read_text()
tpl = re.sub(r"\{\{VALUES:([^}]+)\}\}", lambda m: values_grid(m.group(1)), tpl)
tpl = re.sub(r"\{\{FORMULAS:([^}]+)\}\}", lambda m: formulas_table(m.group(1)), tpl)
tpl = re.sub(r"\{\{EXPERIMENTS:(\d+)\}\}", lambda m: experiments_table(m.group(1)), tpl)
tpl = re.sub(r"\{\{VAL:([^}]+)\}\}", lambda m: show(*m.group(1).split("!")), tpl)
tpl = tpl.replace("{{SPEC}}", full_spec()).replace("{{ASSERTIONS}}", assertions())
tpl = tpl.replace("{{EXPERIMENTS_JSON}}", experiments_json())
n_f = sum(1 for s in spec["sheets"].values() for c in s.values() if "f" in c)
tpl = tpl.replace("{{N_FORMULAS}}", str(n_f)).replace("{{N_ASSERTIONS}}", str(len(spec["assertions"])))
tpl = tpl.replace("{{N_EXPERIMENTS}}", str(len(exps)))
left = re.findall(r"\{\{[^}]+\}\}", tpl)
assert not left, left
(HERE / "Excel_Implementation_Labs_Financial_Analysis.md").write_text(tpl)
print("written", len(tpl), "chars")
