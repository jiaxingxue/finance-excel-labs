"""Build the Excel lab workbook from a single spec, recalc in LibreOffice, verify assertions,
and export a machine-readable JSON spec for a future web app."""
import datetime as dt
import json
import subprocess
import shutil
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

OUT = Path(__file__).parent
D = dt.date

sheets = {}          # sheet -> {addr: value or "=formula"}
fmts = {}            # sheet -> {addr: number_format}
labs = []            # lab metadata + assertions


def put(sheet, addr, value, fmt=None):
    sheets.setdefault(sheet, {})[addr] = value
    if fmt:
        fmts.setdefault(sheet, {})[addr] = fmt


def row(sheet, r, start_col, values, fmt=None):
    for i, v in enumerate(values):
        put(sheet, f"{chr(ord(start_col) + i)}{r}", v, fmt)


# =====================================================================
# Sheet: Map (account mapping)
# =====================================================================
row("Map", 1, "A", ["Account", "Type", "PL_Line", "Sign"])
map_rows = [("Revenue", "Revenue", "Net revenue", 1),
            ("COGS", "Expense", "Cost of sales", -1),
            ("Salaries", "Expense", "Operating expense", -1),
            ("Marketing", "Expense", "Operating expense", -1),
            ("Rent", "Expense", "Operating expense", -1),
            ("Travel", "Expense", "Operating expense", -1)]
for i, r in enumerate(map_rows, start=2):
    row("Map", i, "A", list(r))

# =====================================================================
# Sheet: GL (long-format ledger + budget)
# =====================================================================
row("GL", 1, "A", ["Period", "Dept", "Account", "Version", "Amount"])
gl = []
def add(period, version, account, **by_dept):
    for dept, amt in by_dept.items():
        gl.append((period, dept, account, version, amt))

# 2026-08 budget
add("2026-08", "BUD", "Revenue", Sales=300000, Ops=200000)
add("2026-08", "BUD", "COGS", Ops=200000)
add("2026-08", "BUD", "Salaries", Sales=70000, Ops=50000)
add("2026-08", "BUD", "Marketing", Sales=40000)
add("2026-08", "BUD", "Rent", Ops=25000)
add("2026-08", "BUD", "Travel", Sales=5000, Ops=3000)
# 2026-08 actual
add("2026-08", "ACT", "Revenue", Sales=270000, Ops=198000)
add("2026-08", "ACT", "COGS", Ops=191000)
add("2026-08", "ACT", "Salaries", Sales=72000, Ops=52500)
add("2026-08", "ACT", "Marketing", Sales=52000)
add("2026-08", "ACT", "Rent", Ops=25000)
add("2026-08", "ACT", "Travel", Sales=3600, Ops=1500)
# 2026-07 actual
add("2026-07", "ACT", "Revenue", Sales=262000, Ops=193000)
add("2026-07", "ACT", "COGS", Ops=186000)
add("2026-07", "ACT", "Salaries", Sales=70000, Ops=51000)
add("2026-07", "ACT", "Marketing", Sales=39000)
add("2026-07", "ACT", "Rent", Ops=25000)
add("2026-07", "ACT", "Travel", Sales=4200, Ops=2600)
# 2025-08 actual
add("2025-08", "ACT", "Revenue", Sales=250000, Ops=180000)
add("2025-08", "ACT", "COGS", Ops=176000)
add("2025-08", "ACT", "Salaries", Sales=64000, Ops=46000)
add("2025-08", "ACT", "Marketing", Sales=35000)
add("2025-08", "ACT", "Rent", Ops=23000)
add("2025-08", "ACT", "Travel", Sales=4000, Ops=3000)
for i, r in enumerate(gl, start=2):
    row("GL", i, "A", list(r))
GL_LAST = len(gl) + 1
assert GL_LAST == 37, GL_LAST
R = lambda col: f"GL!${col}$2:${col}${GL_LAST}"   # absolute GL column range

# =====================================================================
# LAB 2: Budget vs Actual (sheet BvA)
# =====================================================================
put("BvA", "A1", "Report period"); put("BvA", "B1", "2026-08")
put("BvA", "A2", "Pct threshold"); put("BvA", "B2", 0.05, "0.0%")
put("BvA", "A3", "Abs threshold"); put("BvA", "B3", 5000, "#,##0")
row("BvA", 5, "A", ["Account", "Type", "Budget", "Actual", "Variance", "Var %",
                    "Fav/Unfav", "Needs comment"])
accts = [r[0] for r in map_rows]
for i, a in enumerate(accts, start=6):
    put("BvA", f"A{i}", a)
    put("BvA", f"B{i}", f"=INDEX(Map!$B$2:$B$7,MATCH($A{i},Map!$A$2:$A$7,0))")
    put("BvA", f"C{i}", f'=SUMIFS({R("E")},{R("C")},$A{i},{R("D")},"BUD",{R("A")},$B$1)', "#,##0")
    put("BvA", f"D{i}", f'=SUMIFS({R("E")},{R("C")},$A{i},{R("D")},"ACT",{R("A")},$B$1)', "#,##0")
put("BvA", "A12", "Operating income"); put("BvA", "B12", "Revenue")
put("BvA", "C12", "=C6-SUM(C7:C11)", "#,##0"); put("BvA", "D12", "=D6-SUM(D7:D11)", "#,##0")
for i in range(6, 13):
    put("BvA", f"E{i}", f"=D{i}-C{i}", "#,##0;(#,##0)")
    put("BvA", f"F{i}", f'=IF(C{i}=0,"n/a",E{i}/ABS(C{i}))', "0.0%")
    put("BvA", f"G{i}", f'=IF(D{i}=C{i},"On budget",IF((B{i}="Revenue")=(D{i}>C{i}),"Favorable","Unfavorable"))')
    put("BvA", f"H{i}", f"=IF(ISNUMBER(F{i}),AND(ABS(E{i})>=$B$3,ABS(F{i})>=$B$2),FALSE)")
put("BvA", "A14", "Check: mapped actuals minus all GL actuals (should be 0)")
put("BvA", "D14", f'=SUM(D6:D11)-SUMIFS({R("E")},{R("D")},"ACT",{R("A")},$B$1)', "#,##0")

# =====================================================================
# LAB 3: Flux analysis (sheet Flux)
# =====================================================================
put("Flux", "A1", "Current period"); put("Flux", "B1", "2026-08")
put("Flux", "A2", "Prior month"); put("Flux", "B2", "2026-07")
put("Flux", "A3", "Prior year"); put("Flux", "B3", "2025-08")
put("Flux", "A4", "Pct threshold"); put("Flux", "B4", 0.10, "0.0%")
put("Flux", "A5", "Abs threshold"); put("Flux", "B5", 5000, "#,##0")
row("Flux", 7, "A", ["Account", "Current", "Prior month", "MoM $", "MoM %",
                     "Prior year", "YoY $", "YoY %", "Flag"])
for i, a in enumerate(accts, start=8):
    put("Flux", f"A{i}", a)
    for col, ref in (("B", "$B$1"), ("C", "$B$2"), ("F", "$B$3")):
        put("Flux", f"{col}{i}", f'=SUMIFS({R("E")},{R("C")},$A{i},{R("D")},"ACT",{R("A")},{ref})', "#,##0")
    put("Flux", f"D{i}", f"=B{i}-C{i}", "#,##0;(#,##0)")
    put("Flux", f"E{i}", f'=IF(C{i}=0,"n/a",D{i}/ABS(C{i}))', "0.0%")
    put("Flux", f"G{i}", f"=B{i}-F{i}", "#,##0;(#,##0)")
    put("Flux", f"H{i}", f'=IF(F{i}=0,"n/a",G{i}/ABS(F{i}))', "0.0%")
    put("Flux", f"I{i}", f'=IF(OR(AND(ABS(D{i})>=$B$5,ABS(N(E{i}))>=$B$4),'
                         f'AND(ABS(G{i})>=$B$5,ABS(N(H{i}))>=$B$4)),"Explain","")')

# =====================================================================
# LAB 4: Price-Volume-Mix (sheet PVM)
# =====================================================================
row("PVM", 4, "A", ["Product", "Bud units", "Bud price", "Act units", "Act price",
                    "Bud rev", "Act rev", "Bud mix %", "Act units @ bud mix",
                    "Price var", "Volume var", "Mix var", "Total var", "Check"])
pvm = [("Basic", 1000, 50, 900, 52), ("Pro", 500, 120, 600, 115), ("Enterprise", 100, 400, 90, 420)]
for i, (p, bu, bp, au, ap) in enumerate(pvm, start=5):
    row("PVM", i, "A", [p, bu, bp, au, ap])
    put("PVM", f"F{i}", f"=B{i}*C{i}", "#,##0")
    put("PVM", f"G{i}", f"=D{i}*E{i}", "#,##0")
    put("PVM", f"H{i}", f"=B{i}/$B$8", "0.00%")
    put("PVM", f"I{i}", f"=$D$8*H{i}", "#,##0.00")
    put("PVM", f"J{i}", f"=(E{i}-C{i})*D{i}", "#,##0.0;(#,##0.0)")
    put("PVM", f"K{i}", f"=(I{i}-B{i})*C{i}", "#,##0.0;(#,##0.0)")
    put("PVM", f"L{i}", f"=(D{i}-I{i})*C{i}", "#,##0.0;(#,##0.0)")
    put("PVM", f"M{i}", f"=G{i}-F{i}", "#,##0;(#,##0)")
    put("PVM", f"N{i}", f"=ROUND(J{i}+K{i}+L{i}-M{i},6)")
put("PVM", "A8", "Total")
for c in "BDFGHIJKLMN":
    put("PVM", f"{c}8", f"=SUM({c}5:{c}7)")
put("PVM", "C8", "=F8/B8", "0.00"); put("PVM", "E8", "=G8/D8", "0.00")
row("PVM", 10, "A", ["Waterfall step", "Amount"])
for i, (lbl, f) in enumerate([("Budget revenue", "=F8"), ("Price", "=J8"), ("Volume", "=K8"),
                              ("Mix", "=L8"), ("Actual revenue", "=G8")], start=11):
    put("PVM", f"A{i}", lbl); put("PVM", f"B{i}", f, "#,##0.0;(#,##0.0)")

# =====================================================================
# LAB 5: Cost variances and flexible budget (sheet CostVar)
# =====================================================================
put("CostVar", "A1", "Part A - Rate / volume (freight)")
for i, (lbl, v) in enumerate([("Budget shipments", 2000), ("Budget rate / shipment", 12),
                              ("Actual shipments", 2300), ("Actual rate / shipment", 12.5)], start=3):
    put("CostVar", f"A{i}", lbl); put("CostVar", f"B{i}", v)
for i, (lbl, f) in enumerate([("Budget cost", "=B3*B4"), ("Actual cost", "=B5*B6"),
                              ("Rate variance", "=(B6-B4)*B5"), ("Volume variance", "=(B5-B3)*B4"),
                              ("Total variance", "=B9-B8"), ("Check (should be 0)", "=B10+B11-B12")], start=8):
    put("CostVar", f"A{i}", lbl); put("CostVar", f"B{i}", f, "#,##0;(#,##0)")
for i in (10, 11, 12):
    put("CostVar", f"C{i}", f'=IF(B{i}>0,"Unfavorable",IF(B{i}<0,"Favorable","None"))')

put("CostVar", "A15", "Part B - Flexible budget")
flex_inputs = [("Budget units", 10000), ("Budget price", 20), ("Budget variable cost / unit", 12),
               ("Budget fixed costs", 50000), ("Actual units", 11000), ("Actual revenue", 214500),
               ("Actual variable costs", 137500), ("Actual fixed costs", 52000)]
for i, (lbl, v) in enumerate(flex_inputs, start=17):
    put("CostVar", f"A{i}", lbl); put("CostVar", f"B{i}", v, "#,##0")
row("CostVar", 26, "A", ["Line", "Static budget", "Flexible budget", "Actual",
                         "Flex-budget variance (Actual - Flex)", "Sales-volume variance (Flex - Static)"])
lines = {
    27: ("Units", "=$B$17", "=$B$21", "=$B$21"),
    28: ("Revenue", "=B27*$B$18", "=C27*$B$18", "=$B$22"),
    29: ("Variable costs", "=B27*$B$19", "=C27*$B$19", "=$B$23"),
    30: ("Contribution margin", "=B28-B29", "=C28-C29", "=D28-D29"),
    31: ("Fixed costs", "=$B$20", "=$B$20", "=$B$24"),
    32: ("Operating income", "=B30-B31", "=C30-C31", "=D30-D31"),
}
for r_, (lbl, b, c, d) in lines.items():
    put("CostVar", f"A{r_}", lbl)
    put("CostVar", f"B{r_}", b, "#,##0"); put("CostVar", f"C{r_}", c, "#,##0"); put("CostVar", f"D{r_}", d, "#,##0")
    put("CostVar", f"E{r_}", f"=D{r_}-C{r_}", "#,##0;(#,##0)")
    put("CostVar", f"F{r_}", f"=C{r_}-B{r_}", "#,##0;(#,##0)")

# =====================================================================
# LAB 6: Pivot-style cross-tab (sheet PivotLab)
# =====================================================================
put("PivotLab", "A2", "Period"); put("PivotLab", "B2", "2026-08")
put("PivotLab", "A3", "Version"); put("PivotLab", "B3", "ACT")
put("PivotLab", "A5", "Dept \\ Account")
for j, a in enumerate(accts):
    put("PivotLab", f"{chr(ord('B') + j)}5", a)
for i, dept in enumerate(["Sales", "Ops"], start=6):
    put("PivotLab", f"A{i}", dept)
    for j in range(6):
        c = chr(ord("B") + j)
        put("PivotLab", f"{c}{i}",
            f'=SUMIFS({R("E")},{R("B")},$A{i},{R("C")},{c}$5,{R("A")},$B$2,{R("D")},$B$3)', "#,##0")
put("PivotLab", "A8", "Total")
for j in range(6):
    c = chr(ord("B") + j)
    put("PivotLab", f"{c}8", f"=SUM({c}6:{c}7)", "#,##0")
put("PivotLab", "A10", "Check: cross-tab total vs BvA actuals (should be 0)")
put("PivotLab", "B10", "=SUM(B8:G8)-SUM(BvA!D6:D11)", "#,##0")

# =====================================================================
# LAB 7: Bank reconciliation (sheets Book, Bank, BankRec)
# =====================================================================
row("Book", 1, "A", ["ID", "Date", "Amount", "Ref", "Status", "Matched bank ID", "Duplicate?"])
book = [("B1", D(2026, 8, 2), 1200.00, "INV1001"), ("B2", D(2026, 8, 5), -350.00, "CHK2201"),
        ("B3", D(2026, 8, 10), 4999.50, "INV1002"), ("B4", D(2026, 8, 28), -2000.00, "CHK2202"),
        ("B5", D(2026, 8, 31), 780.00, "DEP0831")]
row("Bank", 1, "A", ["ID", "Date", "Amount", "Ref", "Status", "Matched book ID"])
bank = [("K1", D(2026, 8, 3), 1200.00, "INV1001"), ("K2", D(2026, 8, 7), -350.00, "CHK2201"),
        ("K3", D(2026, 8, 11), 4995.00, "INV1002"), ("K4", D(2026, 8, 15), -25.00, "FEE")]
for i, r in enumerate(book, start=2):
    row("Book", i, "A", list(r)); fmts.setdefault("Book", {})[f"B{i}"] = "yyyy-mm-dd"
    fmts["Book"][f"C{i}"] = "#,##0.00"
    put("Book", f"F{i}", f'=IFERROR(LOOKUP(2,1/((Bank!$D$2:$D$5=$D{i})*(ABS(Bank!$C$2:$C$5-$C{i})<=BankRec!$B$2)'
                         f'*(ABS(Bank!$B$2:$B$5-$B{i})<=BankRec!$B$3)),Bank!$A$2:$A$5),"")')
    put("Book", f"E{i}", f'=IF(COUNTIFS(Bank!$D$2:$D$5,$D{i},Bank!$C$2:$C$5,$C{i})>0,"Exact",IF(F{i}<>"","Tolerance","Unmatched"))')
    put("Book", f"G{i}", f'=IF(COUNTIFS($D$2:$D$6,$D{i},$C$2:$C$6,$C{i})>1,"DUPLICATE","")')
for i, r in enumerate(bank, start=2):
    row("Bank", i, "A", list(r)); fmts.setdefault("Bank", {})[f"B{i}"] = "yyyy-mm-dd"
    fmts["Bank"][f"C{i}"] = "#,##0.00"
    put("Bank", f"F{i}", f'=IFERROR(LOOKUP(2,1/((Book!$D$2:$D$6=$D{i})*(ABS(Book!$C$2:$C$6-$C{i})<=BankRec!$B$2)'
                         f'*(ABS(Book!$B$2:$B$6-$B{i})<=BankRec!$B$3)),Book!$A$2:$A$6),"")')
    put("Bank", f"E{i}", f'=IF(COUNTIFS(Book!$D$2:$D$6,$D{i},Book!$C$2:$C$6,$C{i})>0,"Exact",IF(F{i}<>"","Tolerance","Unmatched"))')

put("BankRec", "A1", "Bank reconciliation - August 2026")
put("BankRec", "A2", "Amount tolerance"); put("BankRec", "B2", 5, "#,##0.00")
put("BankRec", "A3", "Day tolerance"); put("BankRec", "B3", 3)
put("BankRec", "A5", "Balance per bank statement"); put("BankRec", "B5", 10000, "#,##0.00")
put("BankRec", "A6", "Balance per GL"); put("BankRec", "B6", 8809.50, "#,##0.00")
rec = [
    (8, "BANK SIDE", None),
    (9, "Balance per bank", "=B5"),
    (10, "+ Deposits in transit", '=SUMIFS(Book!$C$2:$C$6,Book!$E$2:$E$6,"Unmatched",Book!$C$2:$C$6,">0")'),
    (11, "- Outstanding checks", '=-SUMIFS(Book!$C$2:$C$6,Book!$E$2:$E$6,"Unmatched",Book!$C$2:$C$6,"<0")'),
    (12, "Adjusted bank balance", "=B9+B10-B11"),
    (13, "BOOK SIDE", None),
    (14, "Balance per GL", "=B6"),
    (15, "+/- Bank-only items (fees, interest)", '=SUMIFS(Bank!$C$2:$C$5,Bank!$E$2:$E$5,"Unmatched")'),
    (16, "- Book errors (book minus bank on tolerance matches)",
         '=SUMIFS(Book!$C$2:$C$6,Book!$E$2:$E$6,"Tolerance")-SUMIFS(Bank!$C$2:$C$5,Bank!$E$2:$E$5,"Tolerance")'),
    (17, "Adjusted book balance", "=B14+B15-B16"),
    (19, "Difference", "=ROUND(B12-B17,2)"),
    (20, "Status", '=IF(B19=0,"Reconciled","Difference: "&TEXT(B19,"#,##0.00"))'),
]
for r_, lbl, f in rec:
    put("BankRec", f"A{r_}", lbl)
    if f:
        put("BankRec", f"B{r_}", f, "#,##0.00")
put("BankRec", "A22", "Journal entry needed (bank fee + short-paid wire)")
put("BankRec", "B22", "=-(B15-B16)", "#,##0.00")

# =====================================================================
# LAB 8: AR aging and subledger-to-GL (sheets AR, ARSummary)
# =====================================================================
row("AR", 1, "A", ["Invoice", "Customer", "Due date", "Open amount", "Days past due", "Bucket"])
put("AR", "H1", "As of"); put("AR", "H2", D(2026, 8, 31), "yyyy-mm-dd")
ar = [("A1", "Acme", D(2026, 5, 15), 3000), ("A2", "Beta", D(2026, 7, 20), 1500),
      ("A3", "Acme", D(2026, 8, 25), 2200), ("A4", "Cyan", D(2026, 9, 10), 900),
      ("A5", "Beta", D(2026, 6, 25), 1200), ("A6", "Cyan", D(2026, 8, 10), 650)]
for i, r in enumerate(ar, start=2):
    row("AR", i, "A", list(r)); fmts.setdefault("AR", {})[f"C{i}"] = "yyyy-mm-dd"
    put("AR", f"E{i}", f"=$H$2-C{i}", "0")
    put("AR", f"F{i}", f'=IF(E{i}<=0,"Current",IF(E{i}<=30,"1-30",IF(E{i}<=60,"31-60",IF(E{i}<=90,"61-90","90+"))))')

buckets = ["Current", "1-30", "31-60", "61-90", "90+"]
reserve = [0.005, 0.02, 0.05, 0.15, 0.50]
row("ARSummary", 3, "A", ["Bucket", "Amount", "% of total", "Reserve %", "Allowance"])
for i, (b, rp) in enumerate(zip(buckets, reserve), start=4):
    put("ARSummary", f"A{i}", b)
    put("ARSummary", f"B{i}", f"=SUMIFS(AR!$D$2:$D$7,AR!$F$2:$F$7,$A{i})", "#,##0")
    put("ARSummary", f"C{i}", f"=B{i}/$B$9", "0.0%")
    put("ARSummary", f"D{i}", rp, "0.0%")
    put("ARSummary", f"E{i}", f"=B{i}*D{i}", "#,##0.00")
put("ARSummary", "A9", "Total"); put("ARSummary", "B9", "=SUM(B4:B8)", "#,##0")
put("ARSummary", "E9", "=SUMPRODUCT(B4:B8,D4:D8)", "#,##0.00")
put("ARSummary", "A11", "Customer \\ Bucket")
for j, b in enumerate(buckets):
    put("ARSummary", f"{chr(ord('B') + j)}11", b)
put("ARSummary", "G11", "Total")
for i, cust in enumerate(["Acme", "Beta", "Cyan"], start=12):
    put("ARSummary", f"A{i}", cust)
    for j in range(5):
        c = chr(ord("B") + j)
        put("ARSummary", f"{c}{i}", f"=SUMIFS(AR!$D$2:$D$7,AR!$B$2:$B$7,$A{i},AR!$F$2:$F$7,{c}$11)", "#,##0")
    put("ARSummary", f"G{i}", f"=SUM(B{i}:F{i})", "#,##0")
put("ARSummary", "A17", "GL control account (1200-AR)"); put("ARSummary", "B17", 9550, "#,##0")
put("ARSummary", "A18", "Subledger total"); put("ARSummary", "B18", "=SUM(AR!$D$2:$D$7)", "#,##0")
put("ARSummary", "A19", "Difference (investigate)"); put("ARSummary", "B19", "=B17-B18", "#,##0")

# =====================================================================
# LAB 9: Forecasting (sheet Forecast)
# =====================================================================
months = [D(2024 + (m // 12), m % 12 + 1, 1) for m in range(30)]
rng = np.random.default_rng(42)
season = np.tile([0.9, 0.92, 1.0, 1.02, 1.05, 1.1, 1.08, 1.06, 1.0, 0.98, 0.95, 1.2], 3)[:30]
sales = ((100_000 + 1_500 * np.arange(30)) * season + rng.normal(0, 2_000, 30)).round(0)

put("Forecast", "A1", "Monthly sales history (Jan-2024 to Jun-2026) and forecasts")
put("Forecast", "A2", "Train periods (t <=)"); put("Forecast", "B2", 24)
row("Forecast", 4, "A", ["Month", "t", "Sales", "MA3", "Month #", "Trend (all)", "Ratio (all)",
                         "Trend (train)", "Ratio (train)"])
FIRST, LAST = 5, 34
for k in range(30):
    r_ = FIRST + k
    put("Forecast", f"A{r_}", months[k], "mmm-yy")
    put("Forecast", f"B{r_}", k + 1)
    put("Forecast", f"C{r_}", float(sales[k]), "#,##0")
    if k >= 2:
        put("Forecast", f"D{r_}", f"=AVERAGE(C{r_-2}:C{r_})", "#,##0")
    put("Forecast", f"E{r_}", f"=MONTH(A{r_})")
    put("Forecast", f"F{r_}", f"=INTERCEPT($C$5:$C$34,$B$5:$B$34)+SLOPE($C$5:$C$34,$B$5:$B$34)*B{r_}", "#,##0")
    put("Forecast", f"G{r_}", f"=C{r_}/F{r_}", "0.0000")
    put("Forecast", f"H{r_}", f"=INTERCEPT($C$5:$C$28,$B$5:$B$28)+SLOPE($C$5:$C$28,$B$5:$B$28)*B{r_}", "#,##0")
    put("Forecast", f"I{r_}", f"=C{r_}/H{r_}", "0.0000")

row("Forecast", 4, "K", ["Month #", "Raw index (all)", "Seasonal index (all)", "Raw index (train)",
                         "Seasonal index (train)"])
for m in range(1, 13):
    r_ = 4 + m
    put("Forecast", f"K{r_}", m)
    put("Forecast", f"L{r_}", f"=AVERAGEIFS($G$5:$G$34,$E$5:$E$34,K{r_})", "0.0000")
    put("Forecast", f"M{r_}", f"=L{r_}/AVERAGE($L$5:$L$16)", "0.0000")
    put("Forecast", f"N{r_}", f'=AVERAGEIFS($I$5:$I$34,$E$5:$E$34,K{r_},$B$5:$B$34,"<="&$B$2)', "0.0000")
    put("Forecast", f"O{r_}", f"=N{r_}/AVERAGE($N$5:$N$16)", "0.0000")
put("Forecast", "K17", "Average"); put("Forecast", "M17", "=AVERAGE(M5:M16)", "0.0000")
put("Forecast", "O17", "=AVERAGE(O5:O16)", "0.0000")
put("Forecast", "K19", "Slope / month"); put("Forecast", "L19", "=SLOPE($C$5:$C$34,$B$5:$B$34)", "#,##0.0")
put("Forecast", "K20", "Intercept"); put("Forecast", "L20", "=INTERCEPT($C$5:$C$34,$B$5:$B$34)", "#,##0.0")
put("Forecast", "K21", "R-squared"); put("Forecast", "L21", "=RSQ($C$5:$C$34,$B$5:$B$34)", "0.000")

row("Forecast", 37, "A", ["Month", "t", "Linear (FORECAST.LINEAR)", "Linear (TREND)", "Trend x season"])
for k in range(6):
    r_ = 38 + k
    put("Forecast", f"A{r_}", "=EDATE($A$34,1)" if k == 0 else f"=EDATE(A{r_-1},1)", "mmm-yy")
    put("Forecast", f"B{r_}", f"=B{r_-1}+1" if k else "=$B$34+1")
    put("Forecast", f"C{r_}", f"=FORECAST.LINEAR(B{r_},$C$5:$C$34,$B$5:$B$34)", "#,##0")
    put("Forecast", f"D{r_}", f"=TREND($C$5:$C$34,$B$5:$B$34,B{r_})", "#,##0")
    put("Forecast", f"E{r_}", f"=C{r_}*INDEX($M$5:$M$16,MONTH(A{r_}))", "#,##0")

row("Forecast", 46, "A", ["Month (test)", "Actual", "Model (train trend x train index)",
                          "Seasonal naive (t-12)", "APE model", "APE naive"])
for k in range(6):
    r_ = 47 + k
    src = 29 + k           # data rows for t = 25..30
    put("Forecast", f"A{r_}", f"=A{src}", "mmm-yy")
    put("Forecast", f"B{r_}", f"=C{src}", "#,##0")
    put("Forecast", f"C{r_}", f"=H{src}*INDEX($O$5:$O$16,E{src})", "#,##0")
    put("Forecast", f"D{r_}", f"=C{src-12}", "#,##0")
    put("Forecast", f"E{r_}", f"=ABS(B{r_}-C{r_})/ABS(B{r_})", "0.00%")
    put("Forecast", f"F{r_}", f"=ABS(B{r_}-D{r_})/ABS(B{r_})", "0.00%")
put("Forecast", "A54", "MAPE (model)"); put("Forecast", "B54", "=AVERAGE(E47:E52)", "0.00%")
put("Forecast", "A55", "Bias (model)"); put("Forecast", "B55", "=(SUM(C47:C52)-SUM(B47:B52))/SUM(B47:B52)", "+0.00%;-0.00%")
put("Forecast", "A56", "MAPE (seasonal naive)"); put("Forecast", "B56", "=AVERAGE(F47:F52)", "0.00%")
put("Forecast", "A57", "Model beats naive?"); put("Forecast", "B57", "=B54<B56")

# =====================================================================
# LAB 10: Driver-based forecast, scenarios, sensitivity (sheet Drivers)
# =====================================================================
put("Drivers", "A1", "Driver-based subscription revenue forecast")
row("Drivers", 3, "A", ["Driver", "Downside", "Base", "Upside", "Active"])
drv = [("New customers / month", 60, 80, 100, "#,##0"), ("Monthly churn", 0.03, 0.02, 0.015, "0.0%"),
       ("ARPU ($ / month)", 90, 95, 98, "#,##0.00")]
for i, (lbl, a, b, c, f) in enumerate(drv, start=4):
    put("Drivers", f"A{i}", lbl)
    for col, v in zip("BCD", (a, b, c)):
        put("Drivers", f"{col}{i}", v, f)
    put("Drivers", f"E{i}", f"=INDEX(B{i}:D{i},$C$8)", f)
put("Drivers", "A8", "Scenario # (1=Down, 2=Base, 3=Up)"); put("Drivers", "C8", 2)
put("Drivers", "A9", "Active scenario name"); put("Drivers", "C9", "=INDEX($B$3:$D$3,$C$8)")
put("Drivers", "A10", "Starting customers"); put("Drivers", "B10", 1200, "#,##0")

put("Drivers", "A13", "Month")
cols = "BCDEFG"
for j, c in enumerate(cols):
    put("Drivers", f"{c}13", D(2026, 10, 1) if j == 0 else f"=EDATE({cols[j-1]}13,1)", "mmm-yy")
    put("Drivers", f"{c}14", "=$B$10" if j == 0 else f"={cols[j-1]}17", "#,##0")
    put("Drivers", f"{c}15", f"=-{c}14*$E$5", "#,##0;(#,##0)")
    put("Drivers", f"{c}16", "=$E$4", "#,##0")
    put("Drivers", f"{c}17", f"={c}14+{c}15+{c}16", "#,##0")
    put("Drivers", f"{c}18", f"={c}17*$E$6", "#,##0")
for r_, lbl in ((14, "Opening customers"), (15, "Churned"), (16, "New"),
                (17, "Closing customers"), (18, "Revenue")):
    put("Drivers", f"A{r_}", lbl)
put("Drivers", "H13", "6-mo total"); put("Drivers", "H18", "=SUM(B18:G18)", "#,##0")

put("Drivers", "A21", "6-mo revenue by scenario (closed form)")
for col in "BCD":
    put("Drivers", f"{col}21",
        f"={col}6*SUMPRODUCT($B$10*(1-{col}5)^$B$25:$G$25+{col}4*(1-(1-{col}5)^$B$25:$G$25)/{col}5)", "#,##0")
put("Drivers", "A25", "Helper: month index k")
for j, c in enumerate(cols):
    put("Drivers", f"{c}25", j + 1)

put("Drivers", "A28", "Churn \\ ARPU")
for j, v in enumerate([85, 90, 95, 100, 105]):
    put("Drivers", f"{chr(ord('B') + j)}28", v)
for i, ch in enumerate([0.01, 0.015, 0.02, 0.025, 0.03], start=29):
    put("Drivers", f"A{i}", ch, "0.0%")
    for j in range(5):
        c = chr(ord("B") + j)
        put("Drivers", f"{c}{i}",
            f"={c}$28*SUMPRODUCT($B$10*(1-$A{i})^$B$25:$G$25+$E$4*(1-(1-$A{i})^$B$25:$G$25)/$A{i})", "#,##0")

# =====================================================================
# LAB 12: Checks dashboard (sheet Checks)
# =====================================================================
row("Checks", 2, "A", ["Check", "Pass?", "Value"])
checks = [
    ("BvA: mapped actuals tie to GL", "=BvA!D14=0", "=BvA!D14"),
    ("PivotLab: cross-tab ties to BvA", "=PivotLab!B10=0", "=PivotLab!B10"),
    ("PVM: price+volume+mix = total", "=ROUND(PVM!N8,6)=0", "=PVM!N8"),
    ("CostVar: rate+volume = total", "=CostVar!B13=0", "=CostVar!B13"),
    ("BankRec: adjusted balances agree", "=BankRec!B19=0", "=BankRec!B19"),
    ("Forecast: seasonal index averages 1", "=ROUND(Forecast!M17,9)=1", "=Forecast!M17"),
    ("Drivers: schedule = closed form", "=ROUND(Drivers!H18-INDEX(Drivers!B21:D21,Drivers!C8),2)=0",
     "=Drivers!H18-INDEX(Drivers!B21:D21,Drivers!C8)"),
]
for i, (lbl, p, v) in enumerate(checks, start=3):
    put("Checks", f"A{i}", lbl); put("Checks", f"B{i}", p); put("Checks", f"C{i}", v)
put("Checks", "A10", "ALL CONTROL CHECKS PASS")
put("Checks", "B10", "=COUNTIF(B3:B9,FALSE)=0")
put("Checks", "A12", "Open item (informational): AR subledger vs GL difference")
put("Checks", "C12", "=ARSummary!B19")

# =====================================================================
# Independent Python expectations (NOT read from the workbook)
# =====================================================================
exp = {}
def e(sheet, addr, val, tol=0.005):
    exp[(sheet, addr)] = (val, tol)

bva = {"Revenue": (500000, 468000), "COGS": (200000, 191000), "Salaries": (120000, 124500),
       "Marketing": (40000, 52000), "Rent": (25000, 25000), "Travel": (8000, 5100)}
for i, a in enumerate(accts, start=6):
    b, act = bva[a]
    e("BvA", f"C{i}", b); e("BvA", f"D{i}", act); e("BvA", f"E{i}", act - b)
e("BvA", "F6", -0.064, 1e-9); e("BvA", "F9", 0.30, 1e-9); e("BvA", "F11", -0.3625, 1e-9)
e("BvA", "G6", "Unfavorable"); e("BvA", "G7", "Favorable"); e("BvA", "G10", "On budget"); e("BvA", "G11", "Favorable")
e("BvA", "H6", True); e("BvA", "H7", False); e("BvA", "H8", False); e("BvA", "H9", True); e("BvA", "H11", False)
e("BvA", "C12", 107000); e("BvA", "D12", 70400); e("BvA", "E12", -36600); e("BvA", "G12", "Unfavorable"); e("BvA", "H12", True)
e("BvA", "D14", 0)

flux = {"Revenue": (468000, 455000, 430000, ""), "COGS": (191000, 186000, 176000, ""),
        "Salaries": (124500, 121000, 110000, "Explain"), "Marketing": (52000, 39000, 35000, "Explain"),
        "Rent": (25000, 25000, 23000, ""), "Travel": (5100, 6800, 7000, "")}
for i, a in enumerate(accts, start=8):
    cur, pm, py, flag = flux[a]
    e("Flux", f"B{i}", cur); e("Flux", f"C{i}", pm); e("Flux", f"F{i}", py)
    e("Flux", f"D{i}", cur - pm); e("Flux", f"G{i}", cur - py); e("Flux", f"I{i}", flag)
e("Flux", "E11", 13000 / 39000, 1e-9); e("Flux", "H10", 14500 / 110000, 1e-9)

pv = {"Basic": (1800, -312.5, -4687.5, -3200), "Pro": (-3000, -375, 12375, 9000),
      "Enterprise": (1800, -250, -3750, -2200)}
for i, (p, *_ ) in enumerate(pvm, start=5):
    pr, vo, mx, tot = pv[p]
    e("PVM", f"J{i}", pr); e("PVM", f"K{i}", vo); e("PVM", f"L{i}", mx); e("PVM", f"M{i}", tot); e("PVM", f"N{i}", 0)
e("PVM", "F8", 150000); e("PVM", "G8", 153600); e("PVM", "J8", 600); e("PVM", "K8", -937.5)
e("PVM", "L8", 3937.5); e("PVM", "M8", 3600); e("PVM", "N8", 0); e("PVM", "I5", 993.75)

e("CostVar", "B8", 24000); e("CostVar", "B9", 28750); e("CostVar", "B10", 1150); e("CostVar", "B11", 3600)
e("CostVar", "B12", 4750); e("CostVar", "B13", 0); e("CostVar", "C10", "Unfavorable")
e("CostVar", "B32", 30000); e("CostVar", "C32", 38000); e("CostVar", "D32", 25000)
e("CostVar", "E28", -5500); e("CostVar", "E29", 5500); e("CostVar", "E31", 2000)
e("CostVar", "E32", -13000); e("CostVar", "F32", 8000)

xt = {("Sales", "Revenue"): 270000, ("Sales", "COGS"): 0, ("Sales", "Salaries"): 72000,
      ("Sales", "Marketing"): 52000, ("Sales", "Rent"): 0, ("Sales", "Travel"): 3600,
      ("Ops", "Revenue"): 198000, ("Ops", "COGS"): 191000, ("Ops", "Salaries"): 52500,
      ("Ops", "Marketing"): 0, ("Ops", "Rent"): 25000, ("Ops", "Travel"): 1500}
for i, dept in enumerate(["Sales", "Ops"], start=6):
    for j, a in enumerate(accts):
        e("PivotLab", f"{chr(ord('B') + j)}{i}", xt[(dept, a)])
e("PivotLab", "B8", 468000); e("PivotLab", "G8", 5100); e("PivotLab", "B10", 0)

for addr, v in {"E2": "Exact", "E3": "Exact", "E4": "Tolerance", "E5": "Unmatched", "E6": "Unmatched",
                "F2": "K1", "F3": "K2", "F4": "K3", "F5": "", "F6": "", "G2": ""}.items():
    e("Book", addr, v)
for addr, v in {"E2": "Exact", "E4": "Tolerance", "E5": "Unmatched", "F4": "B3"}.items():
    e("Bank", addr, v)
for addr, v in {"B10": 780, "B11": 2000, "B12": 8780, "B15": -25, "B16": 4.5, "B17": 8780,
                "B19": 0, "B20": "Reconciled", "B22": 29.5}.items():
    e("BankRec", addr, v)

for addr, v in {"E2": 108, "F2": "90+", "E3": 42, "F3": "31-60", "E4": 6, "F4": "1-30",
                "E5": -10, "F5": "Current", "E6": 67, "F6": "61-90", "E7": 21, "F7": "1-30"}.items():
    e("AR", addr, v)
for addr, v in {"B4": 900, "B5": 2850, "B6": 1500, "B7": 1200, "B8": 3000, "B9": 9450,
                "E9": 1816.5, "G12": 5200, "G13": 2700, "G14": 1550, "B12": 0, "F12": 3000,
                "C12": 2200, "D13": 1500, "E13": 1200, "B14": 900, "C14": 650,
                "B18": 9450, "B19": 100}.items():
    e("ARSummary", addr, v)

# forecasting expectations computed in numpy, independently of the workbook
t = np.arange(1, 31)
b1, b0 = np.polyfit(t, sales, 1)
trend = b0 + b1 * t
ratio = sales / trend
mon = np.array([m.month for m in months])
raw = np.array([ratio[mon == m].mean() for m in range(1, 13)])
si = raw / raw.mean()
tt = t[:24]; tb1, tb0 = np.polyfit(tt, sales[:24], 1)
ttrend = tb0 + tb1 * t
tratio = sales / ttrend
traw = np.array([tratio[(mon == m) & (t <= 24)].mean() for m in range(1, 13)])
tsi = traw / traw.mean()
pred = ttrend[24:] * tsi[mon[24:] - 1]
act = sales[24:]
naive = sales[12:18]
mape_m = np.mean(np.abs(act - pred) / np.abs(act)); mape_n = np.mean(np.abs(act - naive) / np.abs(act))
bias_m = (pred.sum() - act.sum()) / act.sum()
e("Forecast", "L19", b1, 1e-6); e("Forecast", "L20", b0, 1e-6)
for k in range(12):
    e("Forecast", f"M{5+k}", si[k], 1e-9); e("Forecast", f"O{5+k}", tsi[k], 1e-9)
fut_t = np.arange(31, 37)
lin = b0 + b1 * fut_t
fut_mon = [7, 8, 9, 10, 11, 12]
for k in range(6):
    e("Forecast", f"C{38+k}", lin[k], 1e-6); e("Forecast", f"D{38+k}", lin[k], 1e-6)
    e("Forecast", f"E{38+k}", lin[k] * si[fut_mon[k] - 1], 1e-6)
    e("Forecast", f"C{47+k}", pred[k], 1e-6)
e("Forecast", "B54", mape_m, 1e-12); e("Forecast", "B55", bias_m, 1e-12); e("Forecast", "B56", mape_n, 1e-12)
e("Forecast", "B57", True); e("Forecast", "M17", 1.0, 1e-12)

def drv_total(new, ch, arpu, start=1200, n=6):
    c, tot = start, 0.0
    for _ in range(n):
        c = c * (1 - ch) + new; tot += c * arpu
    return tot
e("Drivers", "C9", "Base")
e("Drivers", "H18", drv_total(80, 0.02, 95), 1e-6)
e("Drivers", "B21", drv_total(60, 0.03, 90), 1e-6); e("Drivers", "C21", drv_total(80, 0.02, 95), 1e-6)
e("Drivers", "D21", drv_total(100, 0.015, 98), 1e-6)
e("Drivers", "G17", 1200 * 0.98 ** 6 + 80 * (1 - 0.98 ** 6) / 0.02, 1e-6)
for i, ch in enumerate([0.01, 0.015, 0.02, 0.025, 0.03], start=29):
    for j, ap in enumerate([85, 90, 95, 100, 105]):
        e("Drivers", f"{chr(ord('B') + j)}{i}", drv_total(80, ch, ap), 1e-6)
for i in range(3, 10):
    e("Checks", f"B{i}", True)
e("Checks", "B10", True); e("Checks", "C12", 100)

# =====================================================================
# Write workbook, recalc in LibreOffice, verify
# =====================================================================
ORDER = ["Map", "GL", "BvA", "Flux", "PVM", "CostVar", "PivotLab", "Book", "Bank", "BankRec",
         "AR", "ARSummary", "Forecast", "Drivers", "Checks"]
assert set(ORDER) == set(sheets)
wb = Workbook(); wb.remove(wb.active)
for name in ORDER:
    ws = wb.create_sheet(name)
    for addr, v in sheets[name].items():
        # Newer functions are stored with the _xlfn. prefix in the .xlsx file format
        if isinstance(v, str) and v.startswith("="):
            v = v.replace("FORECAST.LINEAR(", "_xlfn.FORECAST.LINEAR(")
        ws[addr] = v
        if addr in fmts.get(name, {}):
            ws[addr].number_format = fmts[name][addr]
src = OUT / "excel_labs_src.xlsx"
wb.save(src)

calc_dir = OUT / "calc"; shutil.rmtree(calc_dir, ignore_errors=True); calc_dir.mkdir()
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(calc_dir), str(src)],
               check=True, capture_output=True, timeout=180)
calc = load_workbook(calc_dir / "excel_labs_src.xlsx", data_only=True)

fails, errors = [], []
for name in ORDER:
    ws = calc[name]
    for addr, v in sheets[name].items():
        if isinstance(v, str) and v.startswith("="):
            got = ws[addr].value
            if isinstance(got, str) and got.startswith("#") or got is None and addr not in ("F5", "F6", "G2", "G3", "G4", "G5", "G6"):
                if not (name in ("Book", "Bank") and addr[0] in "FG"):
                    if not (name == "Flux" and addr[0] == "I"):
                        errors.append((name, addr, v, got))
for (name, addr), (val, tol) in exp.items():
    got = calc[name][addr].value
    if got is None and val == "":
        continue
    if isinstance(val, bool) or isinstance(val, str):
        ok = got == val
    else:
        ok = got is not None and not isinstance(got, str) and abs(got - val) <= max(tol, abs(val) * 1e-12)
    if not ok:
        fails.append((name, addr, val, got))
print(f"formula cells: {sum(1 for s in sheets.values() for v in s.values() if isinstance(v,str) and v.startswith('='))}")
print(f"assertions: {len(exp)}  failures: {len(fails)}  error cells: {len(errors)}")
for f in fails[:30]: print("FAIL", f)
for f in errors[:30]: print("ERR", f)

# =====================================================================
# Export JSON spec (cells + formats + assertions with engine-verified values)
# =====================================================================
def jv(v):
    if isinstance(v, dt.date):
        return {"date": v.isoformat()}
    if isinstance(v, (np.floating,)):
        return float(v)
    return v

spec = {"workbook": "fin-analysis-excel-labs", "version": "1.0", "sheetOrder": ORDER, "sheets": {}}
for name in ORDER:
    cells = {}
    for addr, v in sheets[name].items():
        c = {"f": v} if isinstance(v, str) and v.startswith("=") else {"v": jv(v)}
        if addr in fmts.get(name, {}):
            c["fmt"] = fmts[name][addr]
        cells[addr] = c
    spec["sheets"][name] = cells
spec["assertions"] = [
    {"sheet": s, "cell": a, "expected": (float(v) if isinstance(v, (float, np.floating)) and not isinstance(v, bool) else v),
     "tolerance": tol}
    for (s, a), (v, tol) in exp.items()]
(OUT / "spec.json").write_text(json.dumps(spec, indent=1, default=float))

# also dump calculated values for doc generation
vals = {name: {addr: calc[name][addr].value for addr in sheets[name]} for name in ORDER}
(OUT / "values.json").write_text(json.dumps(vals, indent=1, default=str))
